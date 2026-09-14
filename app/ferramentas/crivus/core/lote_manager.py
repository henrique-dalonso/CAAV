"""Leitura da planilha de entrada (export do NPJUR) e geração da planilha
de saída do Processamento em Lote. Henrique, 2026-09-14: colunas reais
confirmadas na planilha de amostra (`config/RELATÓRIO PUBLICAÇÕES NÃO
LIDAS.xlsx`) — NPJUR, DATA DA PUBLICAÇÃO, DATA DA IMPORTAÇÃO DA
PUBLICAÇÃO, TEOR PUBLICAÇÃO, STATUS LEITURA PUBLICAÇÃO (a última não é
usada como entrada, só existe no export). As datas vêm como TEXTO
"DD/MM/AAAA" na planilha real (não como célula de data nativa do Excel)
— confirmado inspecionando a amostra.

Henrique, coordenador, 2026-09-14 (correção no dia seguinte ao ar):
"podemos ignorar 100% a segunda data" (DATA DA IMPORTAÇÃO DA
PUBLICAÇÃO) — só a 1ª data (DATA DA PUBLICAÇÃO) decide alguma coisa em
lote_batch.py. A 2ª coluna deixou de ser obrigatória (planilha sem ela
não quebra o upload); se vier presente, ainda é lida e guardada, só
como registro do dado bruto, sem influenciar nada."""

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


# Henrique, 2026-09-03: TEOR vem MUITO inconsistente em formato (de uma
# frase curta a e-mail inteiro) — decisão já tomada de mandar pra IA como
# está, sem limpeza prévia. DATA DA PUBLICAÇÃO é obrigatória desde
# 2026-09-14: alimenta a decisão de "atrasado" em lote_batch.py.
COLUNAS_OBRIGATORIAS = ["NPJUR", "DATA DA PUBLICAÇÃO", "TEOR PUBLICAÇÃO"]


class PlanilhaInvalida(ValueError):
    """Erro claro (nomeando a coluna faltando) em vez de falha silenciosa
    — mostrado direto na tela de upload."""


def _mapear_colunas(linha_cabecalho):
    mapa = {}
    for indice, celula in enumerate(linha_cabecalho):
        if celula is None:
            continue
        nome = str(celula).strip().upper()
        if nome:
            mapa[nome] = indice

    faltando = [coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in mapa]
    if faltando:
        raise PlanilhaInvalida(
            f"A planilha precisa ter a(s) coluna(s): {', '.join(faltando)}."
        )

    return mapa


def _parsear_data(valor):
    """As datas na planilha real vêm como texto "DD/MM/AAAA" — mas trata
    também o caso de a célula já vir como data/datetime nativa do Excel,
    pra não depender de como cada exportação específica formata isso."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()


def ler_linhas_planilha(caminho):
    """Devolve uma lista de dicts, 1 por linha de dado (cabeçalho
    excluído): {"npjur", "data_publicacao", "data_importacao", "teor"}.
    Linha com teor vazio é pulada (não derruba o upload inteiro) — sem
    nenhuma outra limpeza, por decisão já tomada de não "arrumar" o teor."""
    pasta_trabalho = load_workbook(caminho, read_only=True, data_only=True)
    planilha = pasta_trabalho.active

    linhas_planilha = planilha.iter_rows(values_only=True)
    try:
        cabecalho = next(linhas_planilha)
    except StopIteration:
        raise PlanilhaInvalida("A planilha está vazia.")

    mapa = _mapear_colunas(cabecalho)
    indice_npjur = mapa["NPJUR"]
    indice_data_publicacao = mapa["DATA DA PUBLICAÇÃO"]
    indice_data_importacao = mapa.get("DATA DA IMPORTAÇÃO DA PUBLICAÇÃO")  # opcional, ver docstring do módulo
    indice_teor = mapa["TEOR PUBLICAÇÃO"]

    linhas = []
    for linha in linhas_planilha:
        if len(linha) <= indice_teor:
            continue
        teor = linha[indice_teor]
        if teor is None or not str(teor).strip():
            continue

        data_importacao = None
        if indice_data_importacao is not None and len(linha) > indice_data_importacao:
            data_importacao = _parsear_data(linha[indice_data_importacao])

        linhas.append({
            "npjur": str(linha[indice_npjur]).strip() if linha[indice_npjur] is not None else None,
            "data_publicacao": _parsear_data(linha[indice_data_publicacao]),
            "data_importacao": data_importacao,
            "teor": str(teor).strip(),
        })

    pasta_trabalho.close()
    return linhas


def gerar_planilha_saida(analises, caminho_destino):
    """Planilha NOVA (não a original editada), só com
    NPJUR | Nº DO PROCESSO | STATUS | MOTIVO — Henrique, 2026-09-03.
    STATUS/MOTIVO refletem o resultado do PROCESSAMENTO (a IA rodou ou
    não), não o da revisão humana depois — "OK" com MOTIVO vazio pra
    quem terminou (mesmo que ainda aguarde revisão), "ERRO" com o motivo
    real pra quem falhou.

    "ATRASADO" (Henrique, coordenador, 2026-09-14): a linha nunca chegou
    a ir pra IA de propósito — 2+ dias de atraso desde a publicação,
    encaminhada pro time tratar manualmente (ver `eh_atrasado` em
    lote_batch.py)."""
    pasta_trabalho = Workbook()
    planilha = pasta_trabalho.active
    planilha.append(["NPJUR", "Nº DO PROCESSO", "STATUS", "MOTIVO"])

    for analise in analises:
        if analise.status == "erro":
            planilha.append([analise.npjur or "", "", "ERRO", analise.erro_mensagem or ""])
        elif analise.status == "atrasado":
            planilha.append([analise.npjur or "", "", "ATRASADO", analise.erro_mensagem or ""])
        else:
            planilha.append([analise.npjur or "", analise.processo or "", "OK", ""])

    Path(caminho_destino).parent.mkdir(parents=True, exist_ok=True)
    pasta_trabalho.save(caminho_destino)
