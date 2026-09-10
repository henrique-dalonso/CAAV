"""
Gera o template Word (config/templates/condenacao.docx) usado por
Extratus - Condenacao.

Reaproveita o padrao de identificacao/cronologia/parecer/prazo identico a
gerar_template_relatorio.py (mesmo processo bancario, mesma estrutura). A
parte nova e a secao CALCULO DA CONDENACAO: uma tabela nativa do Word de
8 colunas, preenchida via a tecnica de "row-repeat" do docxtpl
({%tr for %} / {%tr endfor %}) - primeira vez que esta tecnica e usada
neste projeto (confirmado por pesquisa antes de escrever este script:
todo o resto do site usa paragrafo repetido, nunca tabela de verdade).
Ver core/pos_processamento_condenacao.py para os campos que a tabela
espera (itens_calculo_processados, totais, subtotal_1, etc. - todos ja
calculados em codigo, nunca pedidos a IA).

A ressalva obrigatoria ("estimativa... nao substitui conferencia
oficial") e texto FIXO aqui, nao um campo do template - garante que
aparece sempre, formatada certo, sem depender da IA lembrar (mesmo
raciocinio do esqueleto de e-mail da Emenda).

Este script so precisa ser rodado de novo se o template for perdido ou
precisar ser reconstruido do zero. Ajustes normais de visual devem ser
feitos abrindo o .docx gerado direto no Word.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.plataforma.paths import PROJECT_ROOT


CAMINHO_SAIDA = str(
    PROJECT_ROOT / "app" / "ferramentas" / "nucleo_relatorios" / "config" / "templates" / "condenacao.docx"
)

CINZA = RGBColor(0x80, 0x80, 0x80)


def adicionar_borda_inferior(paragrafo, cor="808080", tamanho=6):
    p_pr = paragrafo._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(tamanho))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), cor)
    p_borders.append(bottom)
    p_pr.append(p_borders)


def configurar_pagina(documento):
    secao = documento.sections[0]
    secao.page_height = Cm(29.7)
    secao.page_width = Cm(21.0)
    secao.top_margin = Cm(2.5)
    secao.bottom_margin = Cm(2.5)
    secao.left_margin = Cm(2.5)
    secao.right_margin = Cm(2.5)


def configurar_fonte_padrao(documento):
    estilo = documento.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    estilo.paragraph_format.space_after = Pt(0)
    estilo.paragraph_format.line_spacing = 1.0


def campo_rotulo(documento, rotulo, variavel, espaco_depois=4):
    paragrafo = documento.add_paragraph()
    paragrafo.paragraph_format.space_after = Pt(espaco_depois)
    paragrafo.paragraph_format.line_spacing = 1.0

    run_rotulo = paragrafo.add_run(f"{rotulo}: ")
    run_rotulo.bold = True

    paragrafo.add_run("{{ " + variavel + " }}")

    return paragrafo


def titulo_secao(documento, texto):
    paragrafo = documento.add_paragraph()
    paragrafo.paragraph_format.space_before = Pt(12)
    paragrafo.paragraph_format.space_after = Pt(6)

    run = paragrafo.add_run(texto)
    run.bold = True
    run.font.size = Pt(11)

    adicionar_borda_inferior(paragrafo)

    return paragrafo


def _definir_texto_celula(celula, texto, negrito=False, tamanho_fonte=None):
    """Substitui o conteúdo (vazio, por padrão) da célula por um único
    parágrafo com o texto informado — usado tanto pra rótulos fixos
    quanto pra tags Jinja/docxtpl (`{{ variavel }}`, `{%tr for ... %}`),
    já que pro docxtpl processar a tag ela precisa estar no XML como
    texto de parágrafo comum, igual qualquer outro lugar do documento."""
    paragrafo = celula.paragraphs[0]
    for run_antigo in list(paragrafo.runs):
        run_antigo._element.getparent().remove(run_antigo._element)
    run = paragrafo.add_run(texto)
    run.bold = negrito
    if tamanho_fonte:
        run.font.size = Pt(tamanho_fonte)
    return paragrafo


def construir_tabela_calculo(documento):
    """Tabela nativa do Word, 8 colunas, conforme a planilha de débitos
    judiciais do escritório.

    A tag de repetição de linha do docxtpl (`{%tr ... %}`) precisa ficar
    numa linha SÓ DELA — o docxtpl remove o `<w:tr>` inteiro da linha que
    contém a tag e substitui pela marcação Jinja pura (confirmado testando
    localmente: colocar `{%tr for %}` e `{%tr endfor %}` na MESMA linha
    faz o docxtpl descartar tudo entre os dois, quebrando o template — a
    tag de abertura e a de fechamento precisam estar cada uma na sua
    própria linha, "envolvendo" a linha real de dados no meio, que aí sim
    é uma linha comum, sem prefixo "tr", repetida pelo for/endfor que
    sobrou ao redor dela). Layout final: cabeçalho fixo → linha só com
    `{%tr for %}` → linha de dados normal (`{{ item.x }}`, SEM "tr") →
    linha só com `{%tr endfor %}` → linha TOTAIS fixa."""
    colunas = [
        "ITEM", "DESCRIÇÃO", "DATA", "VALOR SINGELO", "VALOR ATUALIZADO",
        "JUROS COMPENSATÓRIOS", "JUROS MORATÓRIOS", "TOTAL",
    ]
    n_colunas = len(colunas)

    tabela = documento.add_table(rows=5, cols=n_colunas)
    tabela.style = "Table Grid"

    for indice, texto_coluna in enumerate(colunas):
        _definir_texto_celula(tabela.rows[0].cells[indice], texto_coluna, negrito=True, tamanho_fonte=9)

    # Linha-marcador de abertura do loop — só a primeira célula carrega a
    # tag, as demais ficam vazias (a linha inteira desaparece na
    # renderização, então não importa o que tem nas outras células).
    _definir_texto_celula(tabela.rows[1].cells[0], "{%tr for item in itens_calculo_processados %}")

    # Linha de dados de verdade — SEM prefixo "tr", célula por célula.
    linha_dados = tabela.rows[2].cells
    _definir_texto_celula(linha_dados[0], "{{ loop.index }}")
    _definir_texto_celula(linha_dados[1], "{{ item.descricao }}")
    _definir_texto_celula(linha_dados[2], "{{ item.data }}")
    _definir_texto_celula(linha_dados[3], "{{ '%.2f'|format(item.valor_singelo) }}")
    _definir_texto_celula(linha_dados[4], "{{ '%.2f'|format(item.valor_atualizado) }}")
    _definir_texto_celula(linha_dados[5], "{{ '%.2f'|format(item.juros_compensatorios) }}")
    _definir_texto_celula(linha_dados[6], "{{ '%.2f'|format(item.juros_moratorios) }}")
    _definir_texto_celula(linha_dados[7], "{{ '%.2f'|format(item.total) }}")

    # Linha-marcador de fechamento do loop — mesma lógica da abertura.
    _definir_texto_celula(tabela.rows[3].cells[0], "{%tr endfor %}")

    linha_totais = tabela.rows[4].cells
    _definir_texto_celula(linha_totais[0], "")
    _definir_texto_celula(linha_totais[1], "TOTAIS", negrito=True)
    _definir_texto_celula(linha_totais[2], "")
    _definir_texto_celula(linha_totais[3], "{{ '%.2f'|format(totais.valor_singelo) }}", negrito=True)
    _definir_texto_celula(linha_totais[4], "{{ '%.2f'|format(totais.valor_atualizado) }}", negrito=True)
    _definir_texto_celula(linha_totais[5], "{{ '%.2f'|format(totais.juros_compensatorios) }}", negrito=True)
    _definir_texto_celula(linha_totais[6], "{{ '%.2f'|format(totais.juros_moratorios) }}", negrito=True)
    _definir_texto_celula(linha_totais[7], "{{ '%.2f'|format(totais.total) }}", negrito=True)

    return tabela


def construir_secao_calculo(documento):
    """Toda a seção fica dentro de `{% if tem_condenacao_liquida %}` —
    mesma técnica já usada pro loop de cronologia (tag Jinja sozinha num
    parágrafo próprio, docxtpl processa o XML inteiro como template, então
    a condicional funciona atravessando parágrafos/tabela normalmente)."""
    documento.add_paragraph("{% if tem_condenacao_liquida %}")

    titulo_secao(documento, "CÁLCULO DA CONDENAÇÃO")
    construir_tabela_calculo(documento)

    documento.add_paragraph("{% if parametro_supletivo_aplicado %}")
    aviso = documento.add_paragraph(
        "Ao menos um parâmetro (índice, taxa ou termo inicial) não constava "
        "expressamente da decisão e foi aplicado o padrão legal supletivo."
    )
    aviso.paragraph_format.space_before = Pt(6)
    for run in aviso.runs:
        run.italic = True
    documento.add_paragraph("{% endif %}")

    campo_rotulo(documento, "Subtotal", "'%.2f'|format(subtotal_1)")
    campo_rotulo(documento, "Honorários advocatícios", "honorarios_percentual")
    campo_rotulo(documento, "  Valor dos honorários", "'%.2f'|format(honorarios_valor)")
    campo_rotulo(documento, "Subtotal", "'%.2f'|format(subtotal_2)")

    documento.add_paragraph("{% if aplica_multa_523 %}")
    campo_rotulo(documento, "Art. 523, §1º, CPC — multa (10%)", "'%.2f'|format(multa_523_valor)")
    campo_rotulo(documento, "Art. 523, §1º, CPC — honorários (10%)", "'%.2f'|format(honorarios_523_valor)")
    documento.add_paragraph("{% else %}")
    aviso_523 = documento.add_paragraph(
        "Multa e honorários do art. 523, §1º, CPC ainda não exigíveis "
        "(processo fora da fase de cumprimento de sentença, ou prazo de "
        "pagamento voluntário ainda em curso)."
    )
    for run in aviso_523.runs:
        run.italic = True
    documento.add_paragraph("{% endif %}")

    total_geral = campo_rotulo(documento, "TOTAL GERAL", "'%.2f'|format(total_geral)")
    for run in total_geral.runs:
        run.bold = True

    recomendacao = documento.add_paragraph()
    recomendacao.paragraph_format.space_before = Pt(8)
    run_recomendacao = recomendacao.add_run("Recomendação: {{ recomendacao_texto }}")
    run_recomendacao.bold = True
    recomendacao.add_run(" — {{ recomendacao_justificativa }}")

    ressalva = documento.add_paragraph(
        "O cálculo acima é uma estimativa para fins de acompanhamento "
        "estratégico e não substitui a conferência pela calculadora "
        "oficial do tribunal ou por contador judicial antes de qualquer "
        "pagamento ou levantamento de valores."
    )
    ressalva.paragraph_format.space_before = Pt(8)
    for run in ressalva.runs:
        run.italic = True

    documento.add_paragraph("{% endif %}")


def construir_template():
    documento = Document()

    configurar_pagina(documento)
    configurar_fonte_padrao(documento)

    titulo = documento.add_paragraph()
    run_titulo = titulo.add_run("RELATÓRIO PROCESSUAL — CONDENAÇÃO")
    run_titulo.bold = True
    run_titulo.font.size = Pt(14)
    titulo.paragraph_format.space_after = Pt(12)

    campo_rotulo(documento, "TIPO DA AÇÃO", "tipo_acao")
    campo_rotulo(documento, "N° PROCESSO", "numero_processo")
    campo_rotulo(documento, "INCIDENTE", "incidente")
    campo_rotulo(documento, "VALOR DA CAUSA", "valor_causa")
    campo_rotulo(documento, "VALOR DA DÍVIDA AJUIZADA", "valor_divida")
    campo_rotulo(documento, "AUTOR", "autor")
    campo_rotulo(documento, "RÉU", "reu")
    campo_rotulo(documento, "BEM", "bem")
    campo_rotulo(documento, "CONTRATO", "contrato")
    campo_rotulo(documento, "COMARCA/TRIBUNAL", "comarca")

    titulo_secao(documento, "CRONOLOGIA PROCESSUAL")
    documento.add_paragraph("{% for evento in cronologia %}")
    paragrafo_evento = documento.add_paragraph()
    paragrafo_evento.paragraph_format.space_after = Pt(6)
    run_data = paragrafo_evento.add_run("{{ evento.data }}")
    run_data.bold = True
    paragrafo_evento.add_run(" – {{ evento.ator }} – {{ evento.descricao }}")
    documento.add_paragraph("{% endfor %}")

    titulo_secao(documento, "PARECER DO ESCRITÓRIO")
    # "{{r ... }}" ativa o RichText do docxtpl - ver comentário
    # equivalente em gerar_template_relatorio.py (mesmo campo "parecer",
    # mesma regra de negrito na recomendação).
    paragrafo_parecer = documento.add_paragraph("{{r parecer }}")
    paragrafo_parecer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragrafo_parecer.paragraph_format.space_after = Pt(8)

    campo_rotulo(documento, "Data publicação/ciência", "data_publicacao")
    campo_rotulo(documento, "Prazo Fatal ED", "prazo_fatal_ed")
    campo_rotulo(documento, "Prazo Fatal", "prazo_fatal")
    campo_rotulo(documento, "Status atual", "status_atual")

    construir_secao_calculo(documento)

    documento.save(CAMINHO_SAIDA)
    print(f"Template salvo em: {CAMINHO_SAIDA}")


if __name__ == "__main__":
    construir_template()
