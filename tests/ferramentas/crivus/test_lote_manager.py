"""lote_manager.py — leitura da planilha de entrada e geração da
planilha de saída do Processamento em Lote. Colunas reais confirmadas
na amostra do NPJUR (config/RELATÓRIO PUBLICAÇÕES NÃO LIDAS.xlsx)."""

from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from app.ferramentas.crivus.core.lote_manager import (
    PlanilhaInvalida,
    gerar_planilha_saida,
    ler_linhas_planilha,
)


def _criar_planilha(tmp_path, cabecalho, linhas, nome="planilha.xlsx"):
    pasta_trabalho = Workbook()
    planilha = pasta_trabalho.active
    planilha.append(cabecalho)
    for linha in linhas:
        planilha.append(linha)
    caminho = tmp_path / nome
    pasta_trabalho.save(caminho)
    return caminho


def test_le_planilha_com_colunas_na_ordem_real(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        ["NPJUR", "DATA DA PUBLICAÇÃO", "DATA DA IMPORTAÇÃO DA PUBLICAÇÃO", "TEOR PUBLICAÇÃO", "STATUS LEITURA PUBLICAÇÃO"],
        [["0119225", "13/08/2013", "14/08/2013", "teor de teste", "NÃO LIDA"]],
    )

    linhas = ler_linhas_planilha(caminho)

    assert len(linhas) == 1
    assert linhas[0]["npjur"] == "0119225"
    assert linhas[0]["data_publicacao"] == date(2013, 8, 13)
    assert linhas[0]["data_importacao"] == date(2013, 8, 14)
    assert linhas[0]["teor"] == "teor de teste"


def test_le_planilha_com_colunas_fora_de_ordem_e_espaco_extra(tmp_path):
    """A planilha é um export externo (NPJUR), fora do nosso controle —
    a leitura não pode depender da ordem das colunas nem de espaço extra
    no cabeçalho."""
    caminho = _criar_planilha(
        tmp_path,
        ["TEOR PUBLICAÇÃO ", " NPJUR", "DATA DA IMPORTAÇÃO DA PUBLICAÇÃO", "DATA DA PUBLICAÇÃO"],
        [["teor embaralhado", "0055555", "11/09/2026", "10/09/2026"]],
    )

    linhas = ler_linhas_planilha(caminho)

    assert len(linhas) == 1
    assert linhas[0]["npjur"] == "0055555"
    assert linhas[0]["teor"] == "teor embaralhado"
    assert linhas[0]["data_publicacao"] == date(2026, 9, 10)
    assert linhas[0]["data_importacao"] == date(2026, 9, 11)


def test_coluna_obrigatoria_faltando_levanta_erro_claro(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        ["NPJUR", "TEOR PUBLICAÇÃO"],  # falta DATA DA PUBLICAÇÃO
        [["0119225", "teor"]],
    )

    with pytest.raises(PlanilhaInvalida) as excecao:
        ler_linhas_planilha(caminho)

    assert "DATA DA PUBLICAÇÃO" in str(excecao.value)


def test_planilha_sem_a_segunda_data_nao_quebra_o_upload(tmp_path):
    """Henrique, coordenador, 2026-09-14: "podemos ignorar 100% a segunda
    data" — DATA DA IMPORTAÇÃO DA PUBLICAÇÃO deixou de ser obrigatória."""
    caminho = _criar_planilha(
        tmp_path,
        ["NPJUR", "DATA DA PUBLICAÇÃO", "TEOR PUBLICAÇÃO"],
        [["0119225", "13/08/2013", "teor de teste"]],
    )

    linhas = ler_linhas_planilha(caminho)

    assert len(linhas) == 1
    assert linhas[0]["data_publicacao"] == date(2013, 8, 13)
    assert linhas[0]["data_importacao"] is None


def test_planilha_vazia_levanta_erro(tmp_path):
    caminho = tmp_path / "vazia.xlsx"
    Workbook().save(caminho)

    with pytest.raises(PlanilhaInvalida):
        ler_linhas_planilha(caminho)


def test_linha_com_teor_vazio_e_pulada_sem_derrubar_upload(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        ["NPJUR", "DATA DA PUBLICAÇÃO", "DATA DA IMPORTAÇÃO DA PUBLICAÇÃO", "TEOR PUBLICAÇÃO"],
        [
            ["0119225", "13/08/2013", "14/08/2013", "teor válido"],
            ["0222222", "13/08/2013", "14/08/2013", ""],
            ["0333333", "13/08/2013", "14/08/2013", None],
        ],
    )

    linhas = ler_linhas_planilha(caminho)

    assert len(linhas) == 1
    assert linhas[0]["npjur"] == "0119225"


def test_gerar_planilha_saida_com_todos_os_status(tmp_path):
    """Henrique, diretoria, 2026-09-15: "OK" virou "LEITURA REALIZADA";
    "atrasado" e "descartado" (motivos diferentes) viram os dois
    STATUS="EXECUTAR MANUALMENTE", diferenciados só pelo MOTIVO."""
    class _AnaliseFake:
        def __init__(self, npjur, processo, status, erro_mensagem=None):
            self.npjur = npjur
            self.processo = processo
            self.status = status
            self.erro_mensagem = erro_mensagem

    analises = [
        _AnaliseFake("0111111", "0011223-45.2024.8.26.0100", "aguardando_revisao"),
        _AnaliseFake("0222222", None, "erro", erro_mensagem="Falha ao analisar: timeout"),
        _AnaliseFake("0333333", None, "atrasado", erro_mensagem="Publicado há 3 dias. Prazo de 2 dias estourado, encaminhado para tratamento manual."),
        _AnaliseFake("0444444", None, "descartado", erro_mensagem="Conteúdo muito curto ou inválido para análise."),
    ]
    destino = tmp_path / "saida" / "resultado.xlsx"

    gerar_planilha_saida(analises, destino)

    pasta_trabalho = load_workbook(destino)
    planilha = pasta_trabalho.active
    linhas = list(planilha.iter_rows(values_only=True))

    # openpyxl grava string vazia como célula em branco, e devolve None
    # ao reler — mesma coisa na prática (célula vazia na planilha final).
    assert linhas[0] == ("NPJUR", "Nº DO PROCESSO", "STATUS", "MOTIVO")
    assert linhas[1] == ("0111111", "0011223-45.2024.8.26.0100", "LEITURA REALIZADA", None)
    assert linhas[2] == ("0222222", None, "ERRO", "Falha ao analisar: timeout")
    assert linhas[3] == ("0333333", None, "EXECUTAR MANUALMENTE", "Publicado há 3 dias. Prazo de 2 dias estourado, encaminhado para tratamento manual.")
    assert linhas[4] == ("0444444", None, "EXECUTAR MANUALMENTE", "Conteúdo muito curto ou inválido para análise.")
