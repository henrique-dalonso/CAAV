"""Cobre um buraco real encontrado em 2026-09-10: `salvar_relatorio_docx`
nunca teve nenhum teste, então nada travava uma regressão de rótulo
duplicado (achado real em produção — ver [[extratus-motor-unificado]]) ou
de campo faltando no molde. Usa um `dados` sintético, sem IA nenhuma."""

from docx import Document

from app.ferramentas.nucleo_relatorios.core.relatorio_manager import salvar_relatorio_docx, texto_para_richtext


def _dados_sinteticos():
    return {
        "tipo_acao": "Busca e Apreensão",
        "numero_processo": "0012345-67.2026.8.16.0030",
        "incidente": "",
        "valor_causa": "R$ 10.000,00",
        "valor_divida": "R$ 9.500,00",
        "autor": "Banco Exemplo S.A.",
        "reu": "Fulano de Tal",
        "bem": "Veículo Fiat Uno 2015, placa ABC1234",
        "contrato": "123456 - 36 parcelas - 1,5% a.m.",
        "comarca": "1ª Vara Cível - Curitiba/PR",
        "cronologia": [
            {"data": "01/01/2026", "ator": "Autor", "descricao": "Distribuição da inicial."},
            {"data": "10/01/2026", "ator": "Juízo", "descricao": "Liminar deferida."},
        ],
        "parecer": "Recomenda-se acompanhar o cumprimento do mandado.",
        "data_publicacao": "15/01/2026",
        "prazo_fatal_ed": "",
        "prazo_fatal": "30/01/2026",
        "status_atual": "Aguardando cumprimento de mandado.",
    }


def _texto_completo(caminho_docx):
    documento = Document(str(caminho_docx))
    return "\n".join(paragrafo.text for paragrafo in documento.paragraphs)


def test_renderiza_sem_erro_e_inclui_o_conteudo_de_cada_campo(tmp_path):
    caminho_saida = tmp_path / "relatorio_teste.docx"
    dados = _dados_sinteticos()

    salvar_relatorio_docx(dados, caminho_saida)

    assert caminho_saida.exists()
    texto = _texto_completo(caminho_saida)
    assert dados["numero_processo"] in texto
    assert dados["parecer"] in texto
    assert dados["cronologia"][0]["descricao"] in texto


def test_cada_rotulo_do_molde_aparece_uma_unica_vez(tmp_path):
    """Trava a regressão real que motivou este arquivo: o prompt antigo
    fazia a IA embutir o rótulo dentro do próprio valor (ex: "Prazo Fatal:
    30/01/2026" dentro do campo `prazo_fatal`), duplicando o rótulo que o
    molde já escreve em negrito."""
    caminho_saida = tmp_path / "relatorio_teste.docx"
    dados = _dados_sinteticos()

    salvar_relatorio_docx(dados, caminho_saida)

    texto = _texto_completo(caminho_saida)
    rotulos = [
        "TIPO DA AÇÃO",
        "N° PROCESSO",
        "VALOR DA CAUSA",
        "VALOR DA DÍVIDA AJUIZADA",
        "AUTOR",
        "RÉU",
        "CONTRATO",
        "COMARCA/TRIBUNAL",
        "Data publicação/ciência",
        # "Prazo Fatal:" com dois-pontos, não "Prazo Fatal" sozinho — sem o
        # dois-pontos, isso também bate como substring dentro do rótulo
        # separado "Prazo Fatal ED:", inflando a contagem por acidente.
        "Prazo Fatal:",
        "Status atual",
    ]
    for rotulo in rotulos:
        assert texto.count(rotulo) == 1, f"Rótulo '{rotulo}' apareceu {texto.count(rotulo)}x, esperado 1x"


def test_rotulo_embutido_no_valor_e_pego_pelo_teste_acima(tmp_path):
    """Prova que o teste anterior de fato pega o bug real (a IA embutindo
    o rótulo dentro do próprio campo, como o prompt antigo induzia) — sem
    isso, o teste acima poderia estar "verde" por acidente."""
    caminho_saida = tmp_path / "relatorio_com_bug.docx"
    dados = _dados_sinteticos()
    dados["prazo_fatal"] = "Prazo Fatal: 30/01/2026"

    salvar_relatorio_docx(dados, caminho_saida)

    texto = _texto_completo(caminho_saida)
    assert texto.count("Prazo Fatal:") == 2


def test_parecer_com_negrito_markdown_vira_negrito_de_verdade_no_word(tmp_path):
    """Achado real (2026-09-10): o prompt pede a recomendação do parecer
    em negrito, mas o campo só aceita texto puro — a IA usava markdown
    ("**assim**"), que aparecia com os asteriscos literais no Word, sem
    negrito nenhum. Agora o código interpreta essa marcação de verdade."""
    caminho_saida = tmp_path / "relatorio_negrito.docx"
    dados = _dados_sinteticos()
    dados["parecer"] = "Texto normal. **Recomenda-se dispensa de recurso.** Mais texto normal."

    salvar_relatorio_docx(dados, caminho_saida)

    documento = Document(str(caminho_saida))
    texto = "\n".join(p.text for p in documento.paragraphs)
    assert "**" not in texto  # nenhum asterisco literal deve sobrar

    runs_negrito = [
        run.text for p in documento.paragraphs for run in p.runs if run.bold
    ]
    assert "Recomenda-se dispensa de recurso." in runs_negrito


def test_parecer_sem_negrito_continua_igual_a_antes(tmp_path):
    """Texto sem nenhum "**" precisa renderizar exatamente igual a antes
    da mudança pra RichText — não pode virar um regressão silenciosa
    pro caso comum (sem negrito nenhum)."""
    caminho_saida = tmp_path / "relatorio_sem_negrito.docx"
    dados = _dados_sinteticos()

    salvar_relatorio_docx(dados, caminho_saida)

    texto = _texto_completo(caminho_saida)
    assert dados["parecer"] in texto


def test_texto_para_richtext_com_multiplos_trechos_em_negrito():
    """A função em si, isolada — mais de um trecho em negrito no mesmo
    texto, intercalado com texto normal."""
    rt = texto_para_richtext("A **B** C **D** E")
    assert "<w:b/>" in rt.xml  # confirma que pelo menos 1 run saiu com negrito

    from docx import Document as _Document
    documento = _Document()
    documento.add_paragraph("{{r x }}")
    from docxtpl import DocxTemplate
    import io
    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    tpl = DocxTemplate(buffer)
    tpl.render({"x": rt})
    buffer_saida = io.BytesIO()
    tpl.save(buffer_saida)
    buffer_saida.seek(0)
    resultado = _Document(buffer_saida)
    runs = [
        (run.text, bool(run.bold))
        for p in resultado.paragraphs for run in p.runs
        if run.text  # docxtpl pode deixar runs vazios nas bordas, sem relevância aqui
    ]
    assert runs == [("A ", False), ("B", True), (" C ", False), ("D", True), (" E", False)]


def test_texto_para_richtext_com_texto_vazio_nao_quebra():
    assert texto_para_richtext("").xml == ""
    assert texto_para_richtext(None).xml == ""
