"""Fecha o mesmo buraco de cobertura já fechado em Relatórios e Emenda:
renderização do .docx testada com dados sintéticos, sem IA. Cobre também
a técnica NOVA usada aqui (nunca usada antes neste projeto): tabela de
Word de tamanho variável via `{%tr for %}`/`{%tr endfor %}` do docxtpl —
não presumir que "compilou" é prova de que funciona, ver descoberta real
registrada em gerar_template_condenacao.py (a tag de repetição precisa
ficar numa linha só dela, não pode dividir a mesma linha do dado)."""

from docx import Document

from app.ferramentas.nucleo_relatorios.core.relatorio_manager import salvar_relatorio_docx
from app.ferramentas.nucleo_relatorios.tipos import REGISTRO_TIPOS


TEMPLATE = REGISTRO_TIPOS["condenacao"].template_docx_path


def _dados_sem_condenacao_liquida():
    return {
        "tipo_acao": "Busca e Apreensão", "numero_processo": "0011223-45.2026.8.16.0001",
        "incidente": "", "valor_causa": "R$ 10.000,00", "valor_divida": "R$ 9.000,00",
        "autor": "Banco Exemplo S.A.", "reu": "Fulano de Tal", "bem": "Fiat Uno 2015",
        "contrato": "123456", "comarca": "1ª Vara Cível - Curitiba/PR",
        "cronologia": [{"data": "01/01/2026", "ator": "Autor", "descricao": "Distribuição."}],
        "parecer": "Aguardando citação.", "data_publicacao": "", "prazo_fatal_ed": "",
        "prazo_fatal": "", "status_atual": "Em andamento.",
        "tem_condenacao_liquida": False,
        "itens_calculo_processados": [], "totais": None, "subtotal_1": None,
        "honorarios_valor": None, "subtotal_2": None, "aplica_multa_523": False,
        "multa_523_valor": None, "honorarios_523_valor": None, "total_geral": None,
        "recomendacao_texto": "", "recomendacao_justificativa": "",
        "honorarios_percentual": "", "parametro_supletivo_aplicado": False,
    }


def _dados_com_condenacao_liquida(**overrides):
    dados = dict(_dados_sem_condenacao_liquida())
    dados.update({
        "tem_condenacao_liquida": True,
        "parametro_supletivo_aplicado": True,
        "itens_calculo_processados": [
            {"descricao": "Capital", "data": "01/01/2026", "valor_singelo": 1000.0,
             "valor_atualizado": 1034.37, "juros_compensatorios": 0.0, "juros_moratorios": 25.0, "total": 1059.37},
            {"descricao": "Custas", "data": "01/01/2026", "valor_singelo": 200.0,
             "valor_atualizado": 206.87, "juros_compensatorios": 0.0, "juros_moratorios": 5.0, "total": 211.87},
        ],
        "totais": {"valor_singelo": 1200.0, "valor_atualizado": 1241.24, "juros_compensatorios": 0.0,
                   "juros_moratorios": 30.0, "total": 1271.24},
        "subtotal_1": 1271.24, "honorarios_percentual": "10%", "honorarios_valor": 127.12,
        "subtotal_2": 1398.36, "aplica_multa_523": True, "multa_523_valor": 139.84,
        "honorarios_523_valor": 139.84, "total_geral": 1678.04,
        "recomendacao_texto": "NÃO IMPUGNAR", "recomendacao_justificativa": "Valores conferem.",
    })
    dados.update(overrides)
    return dados


def _texto_completo(caminho_docx):
    documento = Document(str(caminho_docx))
    return "\n".join(paragrafo.text for paragrafo in documento.paragraphs)


def test_sem_condenacao_liquida_nao_gera_secao_de_calculo(tmp_path):
    caminho_saida = tmp_path / "condenacao_sem_calculo.docx"

    salvar_relatorio_docx(_dados_sem_condenacao_liquida(), caminho_saida, template_docx_path=TEMPLATE)

    documento = Document(str(caminho_saida))
    texto = "\n".join(p.text for p in documento.paragraphs)
    assert "CÁLCULO DA CONDENAÇÃO" not in texto.upper()
    assert len(documento.tables) == 0


def test_com_condenacao_liquida_tabela_renderiza_com_o_numero_certo_de_linhas(tmp_path):
    """Trava a técnica nova (row-repeat do docxtpl): 2 itens sintéticos
    devem virar exatamente 2 linhas de dados na tabela, mais cabeçalho e
    TOTAIS — nem a mais (tag não removida) nem a menos (loop quebrado)."""
    caminho_saida = tmp_path / "condenacao_com_calculo.docx"

    salvar_relatorio_docx(_dados_com_condenacao_liquida(), caminho_saida, template_docx_path=TEMPLATE)

    documento = Document(str(caminho_saida))
    assert len(documento.tables) == 1

    tabela = documento.tables[0]
    linhas = [[celula.text for celula in linha.cells] for linha in tabela.rows]

    assert len(linhas) == 4  # cabeçalho + 2 itens + TOTAIS
    assert linhas[0][0] == "ITEM"
    assert linhas[1][1] == "Capital"
    assert linhas[2][1] == "Custas"
    assert linhas[3][1] == "TOTAIS"
    assert linhas[3][7] == "1271.24"  # total geral da coluna TOTAL


def test_nenhuma_tag_docxtpl_sobrevive_a_renderizacao(tmp_path):
    """Se a técnica de tabela variável quebrar de um jeito sutil (ex:
    tag de repetição virando texto literal em vez de ser processada),
    sobra "{%tr" ou "{{" cru no documento final — trava exatamente esse
    tipo de regressão silenciosa."""
    caminho_saida = tmp_path / "condenacao_sem_tags_soltas.docx"

    salvar_relatorio_docx(_dados_com_condenacao_liquida(), caminho_saida, template_docx_path=TEMPLATE)

    texto = _texto_completo(caminho_saida)
    documento = Document(str(caminho_saida))
    texto_tabelas = "\n".join(
        celula.text for tabela in documento.tables for linha in tabela.rows for celula in linha.cells
    )

    assert "{%" not in texto and "{%" not in texto_tabelas
    assert "{{" not in texto and "{{" not in texto_tabelas


def test_multa_523_some_quando_nao_aplicavel(tmp_path):
    caminho_saida = tmp_path / "condenacao_sem_multa.docx"
    dados = _dados_com_condenacao_liquida(aplica_multa_523=False, multa_523_valor=0.0, honorarios_523_valor=0.0)

    salvar_relatorio_docx(dados, caminho_saida, template_docx_path=TEMPLATE)

    texto = _texto_completo(caminho_saida).lower()
    assert "não exigíveis" in texto
    assert "multa (10%)" not in texto


def test_ressalva_obrigatoria_e_texto_fixo_sempre_presente(tmp_path):
    """A ressalva não é campo da IA (ver docstring do módulo) — precisa
    aparecer sempre, igual, sem depender de nenhum dado do relatório."""
    caminho_saida = tmp_path / "condenacao_ressalva.docx"

    salvar_relatorio_docx(_dados_com_condenacao_liquida(), caminho_saida, template_docx_path=TEMPLATE)

    texto = _texto_completo(caminho_saida)
    assert "não substitui a conferência pela calculadora oficial" in texto
